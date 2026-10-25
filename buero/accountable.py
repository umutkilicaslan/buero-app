"""Import everything from an Accountable export.

Accountable (Settings → Export) produces a folder or ZIP with
  Data.xlsx                    sheets Invoices (one row per invoice line), Expenses, Clients, Transactions
  Original documents/Sales/    <number>.<uuid>.pdf, the invoice PDFs as sent
  Original documents/Purchases/<YYYYMMDD>_<supplier>_<id>.pdf, the expense receipts

What is imported:
- clients (existing ones with the same name are reused, never changed),
- invoices as issued invoices with their original PDF archived unchanged (they are never re-rendered),
  their lines, VAT treatment and, for paid ones, a payment on the payment date,
- expenses with their receipt.
Bank transactions are not imported from here: the bank CSV import is complete and matches payments.

Running the import twice is safe: invoices whose number already exists and expenses already imported
are skipped. Nothing is written until the person confirms the preview.
"""
import hashlib
import re
import zipfile
from collections import OrderedDict
from dataclasses import dataclass, field
from datetime import date, datetime
from decimal import Decimal, InvalidOperation
from pathlib import Path

from sqlalchemy import func, select
from sqlalchemy.orm import Session

from . import bookkeeping, config, money, services
from .models import Client, Expense, Invoice, InvoiceLine, NumberSequence, OtherIncome
from .services import BusinessError, audit

ZERO = Decimal("0")

# Accountable's accounting codes (SKR 03/04 style) and descriptions → our expense categories
CODE_CATEGORY = {"4920": "telecom", "4925": "telecom", "4964": "software", "4945": "education", "4955": "advice",
                 "4957": "advice", "4930": "equipment", "4940": "equipment", "4910": "equipment", "4600": "advertising",
                 "4610": "advertising", "4670": "transport", "4673": "transport", "4660": "travel", "4676": "travel",
                 "4650": "hospitality", "4654": "hospitality", "4360": "fees_insurance", "4380": "fees_insurance",
                 "4970": "bank_fees", "4210": "office_rent", "4800": "repairs", "4805": "repairs"}
WORD_CATEGORY = [("phone", "telecom"), ("internet", "telecom"), ("mobile", "telecom"), ("software", "software"),
                 ("subscription", "software"), ("training", "education"), ("course", "education"),
                 ("education", "education"), ("tax", "advice"), ("accounting", "advice"), ("legal", "advice"),
                 ("office", "equipment"), ("material", "equipment"), ("book", "equipment"), ("equipment", "equipment"),
                 ("advertis", "advertising"), ("marketing", "advertising"), ("website", "advertising"),
                 ("train", "transport"), ("transport", "transport"), ("taxi", "transport"), ("flight", "transport"),
                 ("hotel", "travel"), ("travel", "travel"), ("restaurant", "hospitality"), ("meal", "hospitality"),
                 ("insurance", "fees_insurance"), ("membership", "fees_insurance"), ("bank", "bank_fees"),
                 ("rent", "office_rent"), ("repair", "repairs")]

SALES_PDF = re.compile(r"^(?P<number>.+?)\.[0-9a-f]{8}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{4}-[0-9a-f]{12}\.pdf$", re.I)
PURCHASE_PDF = re.compile(r"^(?P<date>\d{8})_(?P<supplier>.+)_[0-9A-Za-z]+\.(pdf|jpe?g|png|heic)$", re.I)


# ---------- reading the export ----------

def _text(v) -> str:
    return "" if v is None else str(v).strip()


def _num(v) -> Decimal | None:
    if v is None or _text(v) in ("", "–", "-"):
        return None
    if isinstance(v, (int, float)):
        return money.round_cents(Decimal(str(v)))
    t = _text(v).replace("€", "").replace("%", "").replace(" ", "")
    if "," in t:
        t = t.replace(".", "").replace(",", ".")
    try:
        return money.round_cents(Decimal(t))
    except InvalidOperation:
        raise BusinessError(f"'{v}' is not a number.") from None


def _date(v) -> date | None:
    if v is None or _text(v) == "":
        return None
    if isinstance(v, datetime):
        return v.date()
    if isinstance(v, date):
        return v
    t = _text(v)
    for fmt in ("%d.%m.%Y", "%Y-%m-%d", "%d/%m/%Y"):
        try:
            return datetime.strptime(t, fmt).date()
        except ValueError:
            pass
    raise BusinessError(f"'{v}' is not a date.")


def _norm(name: str) -> str:
    return re.sub(r"[^a-z0-9äöüß]+", "", name.lower())


@dataclass
class Export:
    invoices: list[dict]
    expenses: list[dict]
    clients: list[dict]
    sales_pdfs: dict[str, Path]  # invoice number → file
    purchase_files: list[Path]


def read(folder: Path) -> Export:
    """Read an unpacked export. Files are found by name anywhere in the folder."""
    import openpyxl
    files = [p for p in folder.rglob("*") if p.is_file() and not p.name.startswith(("._", "."))]
    xlsx = [p for p in files if p.name.lower() == "data.xlsx"]
    if not xlsx:
        raise BusinessError("Data.xlsx was not found. Choose the whole Accountable export (the folder or the ZIP).")
    wb = openpyxl.load_workbook(xlsx[0], read_only=True, data_only=True)

    def sheet(name: str) -> list[dict]:
        if name not in wb.sheetnames:
            return []
        rows = wb[name].iter_rows(values_only=True)
        header = [_text(h) for h in next(rows, [])]
        return [dict(zip(header, r)) for r in rows if any(v is not None and _text(v) != "" for v in r)]

    if "Invoices" not in wb.sheetnames and "Expenses" not in wb.sheetnames:
        raise BusinessError("This Data.xlsx has no 'Invoices' or 'Expenses' sheet. Is it an Accountable export?")
    sales, purchases = {}, []
    for p in files:
        if m := SALES_PDF.match(p.name):
            sales[m["number"]] = p
        elif PURCHASE_PDF.match(p.name):
            purchases.append(p)
    return Export(sheet("Invoices"), sheet("Expenses"), sheet("Clients"), sales, sorted(purchases))


def unpack(content: bytes, target: Path) -> Path:
    """Unpack an uploaded ZIP safely (no paths outside the target)."""
    import io
    try:
        zf = zipfile.ZipFile(io.BytesIO(content))
    except zipfile.BadZipFile:
        raise BusinessError("This is not a ZIP file.") from None
    for info in zf.infolist():
        name = Path(info.filename)
        if info.is_dir() or name.is_absolute() or ".." in name.parts or name.name.startswith("._"):
            continue
        out = target / name
        out.parent.mkdir(parents=True, exist_ok=True)
        out.write_bytes(zf.read(info))
    return target


# ---------- planning ----------

@dataclass
class PlannedInvoice:
    number: str
    client_name: str
    issue_date: date | None
    due_date: date | None
    paid_date: date | None
    status: str  # issued | paid
    treatment: str
    rate: Decimal
    lines: list[dict]
    net: Decimal
    vat: Decimal
    gross: Decimal
    pdf: Path | None
    skip: str = ""
    notes: list[str] = field(default_factory=list)


@dataclass
class PlannedExpense:
    date: date
    vendor: str
    gross: Decimal
    vat: Decimal
    category: str
    description: str
    business_share: Decimal
    useful_life_years: int | None
    receipt: Path | None
    skip: str = ""
    notes: list[str] = field(default_factory=list)


@dataclass
class Plan:
    new_clients: list[dict]
    invoices: list[PlannedInvoice]
    expenses: list[PlannedExpense]
    unmatched_receipts: list[str]
    warnings: list[str]

    @property
    def invoices_to_import(self):
        return [i for i in self.invoices if not i.skip]

    @property
    def expenses_to_import(self):
        return [e for e in self.expenses if not e.skip]


def _treatment(vat_type: str, vat: Decimal) -> tuple[str, list[str]]:
    t = vat_type.lower()
    if "franchise" in t or "kleinunternehmer" in t or "small business" in t:
        return "kleinunternehmer", []
    if "reverse" in t or "intra" in t:
        return "reverse_charge", []
    if "outside" in t or "export" in t or "non-eu" in t:
        return "non_eu", []
    if vat > 0 or t in ("normal", "standard", ""):
        return "domestic", []
    return "exempt_4_21", [f"Accountable VAT type '{vat_type}' was imported as tax-exempt (§ 4 UStG). "
                           "Check the exemption paragraph on the original PDF."]


def _number_parts(number: str) -> tuple[int, int] | None:
    if m := re.fullmatch(r"(\d{4})-(\d+)", number):
        return int(m[1]), int(m[2])
    if m := re.fullmatch(r"(\d+)-(\d{4})", number):
        return int(m[2]), int(m[1])
    return None


def _category(code: str, description: str) -> tuple[str, list[str]]:
    if code in CODE_CATEGORY:
        return CODE_CATEGORY[code], []
    d = description.lower()
    for word, cat in WORD_CATEGORY:
        if word in d:
            return cat, []
    return "other", [f"Category '{description or code}' is unknown here and was booked as 'Other'. Check it."]


def plan(s: Session, ex: Export) -> Plan:
    warnings: list[str] = []
    existing_clients = {_norm(c.name): c for c in s.scalars(select(Client))}
    sheet_clients = {_norm(_text(c.get("Client Name"))): c for c in ex.clients if _text(c.get("Client Name"))}
    new_clients = []
    for key, row in sheet_clients.items():
        if key not in existing_clients:
            new_clients.append(row)

    # invoices: one row per line, grouped by number in export order
    grouped: "OrderedDict[str, list[dict]]" = OrderedDict()
    for r in ex.invoices:
        n = _text(r.get("Revenue Number"))
        if n:
            grouped.setdefault(n, []).append(r)
    taken = {n for (n,) in s.execute(select(Invoice.number).where(Invoice.number.is_not(None)))}
    invoices = []
    for number, rows in grouped.items():
        first = rows[0]
        kind = _text(first.get("Invoice Type")).lower()
        status = _text(first.get("Status")).lower()
        lines, net, vat, gross = [], ZERO, ZERO, ZERO
        for i, r in enumerate(rows, start=1):
            line_net = _num(r.get("Amount Excluding VAT")) or ZERO
            line_vat = _num(r.get("VAT Amount")) or ZERO
            line_gross = _num(r.get("Total Amount"))
            line_gross = line_net + line_vat if line_gross is None else line_gross
            qty = _num(r.get("Quantity")) or Decimal("1")
            price = _num(r.get("Unit Price"))
            lines.append({"description": _text(r.get("Item Name")) or "Leistung", "quantity": qty,
                          "unit_price": price if price is not None else line_net, "line_net": line_net})
            net, vat, gross = net + line_net, vat + line_vat, gross + line_gross
        treatment, notes = _treatment(_text(first.get("VAT Type")), vat)
        rate = _num(first.get("VAT Rate")) or (money.round_cents(vat * 100 / net) if net and vat else ZERO)
        pi = PlannedInvoice(number=number, client_name=_text(first.get("Client Name")) or "Unknown client",
                            issue_date=_date(first.get("Invoice Date")), due_date=_date(first.get("Due Date")),
                            paid_date=_date(first.get("Payment Date")), status="issued", treatment=treatment,
                            rate=rate, lines=lines, net=net, vat=vat, gross=gross, pdf=ex.sales_pdfs.get(number),
                            notes=notes)
        if number in taken:
            pi.skip = "already in Büro"
        elif "credit" in kind or gross < 0:
            pi.skip = "credit note: add it by hand (Storno)"
        elif status in ("draft", "concept"):
            pi.skip = "draft in Accountable"
        elif pi.issue_date is None:
            pi.skip = "no invoice date"
        elif gross <= 0:
            pi.skip = "amount is 0"
        else:
            if status == "paid":
                if pi.paid_date:
                    pi.status = "paid"
                else:
                    pi.notes.append("Marked paid in Accountable but without a payment date: imported as open.")
            elif "partial" in status:
                pi.notes.append("Part-paid in Accountable: imported as open, record the payments you received.")
            if pi.pdf is None:
                pi.notes.append("No PDF in the export for this invoice.")
        invoices.append(pi)

    # expenses
    have = {(e.date, _norm(e.vendor), e.gross_eur) for e in s.scalars(select(Expense).where(Expense.voided.is_(False)))}
    receipts_left = list(ex.purchase_files)
    expenses = []
    for r in ex.expenses:
        vendor = _text(r.get("Supplier Name")) or "Unknown supplier"
        booking = _date(r.get("Booking Date"))
        paid = _date(r.get("Payment Date"))
        gross = _num(r.get("Total Amount")) or ZERO
        cat, notes = _category(_text(r.get("Accounting Code")), _text(r.get("Code Description")))
        months = _num(r.get("Amortizated Over X Months"))
        life = None
        if months and months > 12:
            life = max(1, int((months / 12).to_integral_value()))
            cat = "asset"
        share = _num(r.get("Professional Use"))
        share = Decimal("100") if share is None else (share * 100 if share <= 1 else share)
        receipt = None
        if booking:
            stem = f"{booking:%Y%m%d}_"
            for p in receipts_left:
                m = PURCHASE_PDF.match(p.name)
                if p.name.startswith(stem) and _norm(m["supplier"]) == _norm(vendor):
                    receipt = p
                    break
            if receipt:
                receipts_left.remove(receipt)
        pe = PlannedExpense(date=paid or booking, vendor=vendor, gross=gross, vat=_num(r.get("VAT Amount")) or ZERO,
                            category=cat, description=_text(r.get("Code Description")), business_share=share,
                            useful_life_years=life, receipt=receipt, notes=notes)
        if pe.date is None:
            pe.skip = "no date"
        elif _text(r.get("Credit Note")).lower() in ("yes", "ja", "true", "1") or gross < 0:
            pe.skip = "credit note: add it by hand"
        elif gross == 0:
            pe.skip = "amount is 0"
        elif _text(r.get("Paid")).lower() in ("no", "nein", "false", "0") and not paid:
            pe.skip = "not paid yet in Accountable (the EÜR counts paid expenses)"
        elif (pe.date, _norm(vendor), gross) in have:
            pe.skip = "already in Büro"
        elif receipt is None and _text(r.get("Document Attached")).lower() == "yes":
            pe.notes.append("Receipt not found in the export.")
        expenses.append(pe)

    # payments that could be counted twice
    paid_dates = [i.paid_date for i in invoices if not i.skip and i.status == "paid"]
    if paid_dates:
        start, end = min(paid_dates), max(paid_dates)
        n, total = s.execute(select(func.count(), func.coalesce(func.sum(OtherIncome.gross_eur), 0))
                             .where(OtherIncome.voided.is_(False), OtherIncome.date >= start,
                                    OtherIncome.date <= end)).one()
        if n:
            warnings.append(f"You already have {n} 'other income' entries ({money.fmt_eur(Decimal(str(total)))}) "
                            f"between {start:%d.%m.%Y} and {end:%d.%m.%Y}. If they are payments for these invoices, "
                            "the money would count twice. Then void them under Income after the import, "
                            "or import without payments.")
    unmatched = [p.name for p in receipts_left]
    if unmatched:
        warnings.append(f"{len(unmatched)} receipt files in the export have no expense row in Data.xlsx and were "
                        "not imported. Add them as expenses by hand if they are business costs.")
    return Plan(new_clients, invoices, expenses, unmatched, warnings)


# ---------- applying ----------

def _client(s: Session, name: str, sheet_row: dict | None, cache: dict) -> Client:
    key = _norm(name)
    if key in cache:
        return cache[key]
    c = s.scalar(select(Client).where(func.lower(Client.name) == name.lower()))
    if c is None:
        c = next((x for x in s.scalars(select(Client)) if _norm(x.name) == key), None)
    if c is None:
        row = sheet_row or {}
        country = (_text(row.get("Country")) or "DE").upper()[:2]
        c = Client(name=name, street=_text(row.get("Street Name And Number")), postal_code=_text(row.get("Post Code")),
                   city=_text(row.get("City")), country=country, vat_id=_text(row.get("VAT Number")),
                   email=_text(row.get("E-mail")), is_business=_text(row.get("Client Type")).lower() != "individual",
                   notes="Imported from Accountable.")
        s.add(c)
        s.flush()
        audit(s, "client", c.id, "imported", source="accountable", name=name)
    cache[key] = c
    return c


def apply(s: Session, ex: Export, p: Plan, record_payments: bool = True) -> dict:
    st = services.get_settings(s)
    seller = services.seller_snapshot(st)
    import json
    sheet_clients = {_norm(_text(c.get("Client Name"))): c for c in ex.clients}
    cache: dict = {}
    counts = {"clients": 0, "invoices": 0, "payments": 0, "expenses": 0, "receipts": 0}
    before = s.scalar(select(func.count()).select_from(Client))
    max_seq: dict[int, int] = {}
    for pi in p.invoices_to_import:
        client = _client(s, pi.client_name, sheet_clients.get(_norm(pi.client_name)), cache)
        inv = Invoice(kind="invoice", status="draft", client=client, language="de", source="accountable",
                      issue_date=pi.issue_date, due_date=pi.due_date)
        s.add(inv)
        s.flush()
        for pos, ln in enumerate(pi.lines, start=1):
            inv.lines.append(InvoiceLine(position=pos, description=ln["description"], quantity=ln["quantity"],
                                         unit="", unit_price=ln["unit_price"], line_net=ln["line_net"]))
        s.flush()
        parts = _number_parts(pi.number)
        inv.number, inv.year = pi.number, pi.issue_date.year
        if parts:
            max_seq[parts[0]] = max(max_seq.get(parts[0], 0), parts[1])
        inv.vat_treatment, inv.vat_rate = pi.treatment, pi.rate
        inv.vat_note = services.vat.NOTES[pi.treatment]["de"] if hasattr(services, "vat") else ""
        inv.net_total, inv.vat_total, inv.gross_total = pi.net, pi.vat, pi.gross
        inv.seller_snapshot = json.dumps(seller, ensure_ascii=False)
        inv.client_snapshot = json.dumps(services.client_snapshot(client, bool(client.vat_id)), ensure_ascii=False)
        inv.issued_at = datetime.combine(pi.issue_date, datetime.min.time())
        inv.status = "issued"
        s.flush()
        audit(s, "invoice", inv.id, "imported", source="accountable", number=pi.number, gross=pi.gross,
              notes=pi.notes)
        if pi.pdf is not None:
            content = pi.pdf.read_bytes()
            target = config.invoices_dir(inv.year) / services.pdf_filename(inv)
            if target.exists() and hashlib.sha256(target.read_bytes()).hexdigest() != hashlib.sha256(content).hexdigest():
                target = target.with_name(f"{target.stem}_accountable{target.suffix}")
            services.attach_pdf(s, inv, content, target)
        counts["invoices"] += 1
        if record_payments and pi.status == "paid":
            services.mark_paid(s, inv, pi.paid_date, pi.gross)
            counts["payments"] += 1

    # keep numbering gap-free: new invoices continue after the highest imported number of each year
    for year, seq in max_seq.items():
        row = s.get(NumberSequence, year) or NumberSequence(year=year, next_number=1)
        if row.next_number <= seq:
            old = row.next_number
            row.next_number = seq + 1
            s.add(row)
            audit(s, "number_sequence", year, "continued_after_import", old=old, new=seq + 1)

    for pe in p.expenses_to_import:
        e = bookkeeping.save_expense(s, None, {
            "date": pe.date, "vendor": pe.vendor, "description": pe.description, "category": pe.category,
            "payment_method": "bank", "currency": "EUR", "gross_eur": pe.gross, "vat_eur": pe.vat,
            "vat_rate": money.round_cents(pe.vat * 100 / (pe.gross - pe.vat)) if pe.vat else ZERO,
            "business_share": pe.business_share, "useful_life_years": pe.useful_life_years,
            "source": "accountable"})
        counts["expenses"] += 1
        if pe.receipt is not None:
            bookkeeping.store_receipt(s, e, pe.receipt.name, pe.receipt.read_bytes())
            counts["receipts"] += 1
    counts["clients"] = s.scalar(select(func.count()).select_from(Client)) - before
    audit(s, "import", None, "accountable", **counts, record_payments=record_payments)
    return counts
